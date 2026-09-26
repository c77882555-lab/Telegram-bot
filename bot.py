import os
import logging
import asyncio
import re
from pathlib import Path

from telegram import Update
from telegram.ext import (
    Application,
    CommandHandler,
    MessageHandler,
    ContextTypes,
    filters,
)

from google import genai
from pypdf import PdfReader
from docx import Document


# =========================
# إعدادات التسجيل
# =========================

logging.basicConfig(
    format="%(asctime)s - %(levelname)s - %(message)s",
    level=logging.INFO,
)
logger = logging.getLogger(__name__)


# =========================
# مفاتيح التشغيل
# =========================

TELEGRAM_TOKEN = os.environ.get("TELEGRAM_TOKEN")
GEMINI_API_KEY = os.environ.get("GEMINI_API_KEY")

if not TELEGRAM_TOKEN:
    raise RuntimeError("TELEGRAM_TOKEN is missing")

if not GEMINI_API_KEY:
    raise RuntimeError("GEMINI_API_KEY is missing")

client = genai.Client(api_key=GEMINI_API_KEY)

MODEL_NAME = "gemini-2.5-flash"

# حجم كل جزء من النص بالأحرف
CHUNK_SIZE = 8000


# =========================
# رسالة البداية
# =========================

async def start(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
) -> None:
    await update.message.reply_text(
        "مرحباً بك يا سجاد! 📚\n\n"
        "أرسل ملف PDF يحتوي على محاضرتك، "
        "وسأترجمه إلى العربية الفصحى بأسلوب أكاديمي "
        "وأرسل لك الترجمة كاملة بملف Word.\n\n"
        "يمكنك إرسال محاضرات طويلة تتجاوز 3000 كلمة."
    )


# =========================
# تقسيم النص إلى أجزاء
# =========================

def split_text(text: str, max_chars: int = CHUNK_SIZE) -> list[str]:
    """
    تقسيم النص إلى أجزاء مع محاولة الحفاظ على الفقرات.
    إذا كانت الفقرة طويلة جداً، يتم تقسيمها إلى أجزاء أصغر.
    """

    text = text.strip()

    if not text:
        return []

    paragraphs = text.split("\n")
    chunks = []
    current_chunk = ""

    for paragraph in paragraphs:
        paragraph = paragraph.strip()

        if not paragraph:
            continue

        # تقسيم الفقرة الطويلة إذا تجاوزت الحد
        if len(paragraph) > max_chars:
            if current_chunk:
                chunks.append(current_chunk)
                current_chunk = ""

            for i in range(0, len(paragraph), max_chars):
                chunks.append(paragraph[i:i + max_chars])

            continue

        candidate = (
            current_chunk + "\n\n" + paragraph
            if current_chunk
            else paragraph
        )

        if len(candidate) <= max_chars:
            current_chunk = candidate
        else:
            if current_chunk:
                chunks.append(current_chunk)

            current_chunk = paragraph

    if current_chunk:
        chunks.append(current_chunk)

    return chunks


# =========================
# ترجمة جزء واحد باستخدام Gemini
# =========================

def translate_chunk(text: str, part_number: int, total_parts: int) -> str:
    if not text.strip():
        return ""

    prompt = f"""
أنت مترجم أكاديمي متخصص في ترجمة المحاضرات الجامعية،
خصوصاً العلوم الطبية والتحليلات المرضية.

المطلوب:
- ترجم النص الإنجليزي التالي إلى العربية الفصحى بدقة علمية.
- حافظ على المعنى العلمي والمصطلحات الطبية.
- اكتب المصطلح الإنجليزي بين قوسين عند الحاجة لتوضيح المصطلحات المهمة.
- حافظ على الأرقام والعناوين والقوائم والاختصارات العلمية.
- لا تختصر النص ولا تلخصه ولا تحذف أي معلومة.
- لا تضف شرحاً من عندك.
- أخرج الترجمة فقط دون مقدمة أو خاتمة.
- هذا الجزء رقم {part_number} من أصل {total_parts}.
- حافظ على ترابط النص مع الأجزاء الأخرى.

النص المطلوب ترجمته:

{text}
"""

    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=prompt,
    )

    translated = response.text

    if not translated or not translated.strip():
        raise ValueError("Gemini returned an empty translation")

    return translated.strip()


# =========================
# ترجمة النص الكامل على أجزاء
# =========================

async def translate_full_text(text: str) -> str:
    chunks = split_text(text)

    if not chunks:
        return ""

    translated_parts = []

    total_parts = len(chunks)

    for index, chunk in enumerate(chunks, start=1):
        logger.info(
            "Translating part %s of %s",
            index,
            total_parts,
        )

        translated = await asyncio.to_thread(
            translate_chunk,
            chunk,
            index,
            total_parts,
        )

        translated_parts.append(translated)

    return "\n\n".join(translated_parts)


# =========================
# معالجة ملفات PDF
# =========================

async def handle_document(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
) -> None:

    message = update.message
    document = message.document

    if not document:
        return

    file_name = document.file_name or "lecture.pdf"

    if not file_name.lower().endswith(".pdf"):
        await message.reply_text(
            "❌ حالياً البوت يدعم ملفات PDF فقط.\n"
            "أرسل المحاضرة بصيغة PDF."
        )
        return

    safe_name = Path(file_name).stem
    safe_name = re.sub(r"[^\w\-]+", "_", safe_name)
    safe_name = safe_name[:80] or "lecture"

    local_input = f"input_{document.file_unique_id}.pdf"
    output_docx = f"Translated_{document.file_unique_id}.docx"

    try:
        await message.reply_text(
            "⏳ استلمت المحاضرة!\n\n"
            "جاري قراءة الملف وتقسيم النص إلى أجزاء "
            "وترجمته بالكامل.\n"
            "قد تستغرق العملية بعض الوقت إذا كانت المحاضرة طويلة."
        )

        telegram_file = await document.get_file()
        await telegram_file.download_to_drive(local_input)

        # قراءة النص من PDF
        reader = PdfReader(local_input)

        full_text = ""

        for page_number, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text()

            if page_text:
                full_text += page_text + "\n\n"

        if not full_text.strip():
            await message.reply_text(
                "❌ لم أتمكن من استخراج النص من الملف.\n"
                "قد يكون الملف عبارة عن صور ممسوحة ضوئياً "
                "ويحتاج إلى تقنية OCR."
            )
            return

        word_count = len(full_text.split())

        await message.reply_text(
            f"📄 تم استخراج النص بنجاح.\n"
            f"عدد الكلمات التقريبي: {word_count}\n\n"
            f"🔄 جاري ترجمة المحاضرة كاملة..."
        )

        # ترجمة جميع الأجزاء
        translated_full = await translate_full_text(full_text)

        if not translated_full.strip():
            await message.reply_text(
                "❌ لم يتم إنشاء ترجمة. حاول مرة أخرى."
            )
            return

        # إنشاء ملف Word
        output_doc = Document()

        output_doc.add_heading(
            f"ترجمة المحاضرة: {safe_name}",
            level=1,
        )

        output_doc.add_paragraph(
            f"عدد كلمات النص الأصلي: {word_count}"
        )

        for paragraph in translated_full.split("\n"):
            if paragraph.strip():
                output_doc.add_paragraph(paragraph.strip())

        output_doc.save(output_docx)

        # إرسال ملف الترجمة
        with open(output_docx, "rb") as translated_file:
            await message.reply_document(
                document=translated_file,
                filename=f"Translated_{safe_name}.docx",
                caption="✅ تمت ترجمة المحاضرة وإعداد ملف Word."
            )

    except Exception as e:
        logger.exception("Error processing document: %s", e)

        await message.reply_text(
            "❌ حدث خطأ أثناء معالجة المحاضرة.\n"
            "تأكد من أن الملف سليم وأن خدمة الترجمة تعمل، "
            "ثم حاول مرة أخرى."
        )

    finally:
        for path in (local_input, output_docx):
            if os.path.exists(path):
                try:
                    os.remove(path)
                except OSError:
                    logger.warning("Could not remove temporary file: %s", path)


# =========================
# تشغيل البوت
# =========================

def main() -> None:
    application = Application.builder().token(
        TELEGRAM_TOKEN
    ).build()

    application.add_handler(
        CommandHandler("start", start)
    )

    application.add_handler(
        MessageHandler(
            filters.Document.ALL,
            handle_document,
        )
    )

    logger.info("Bot is running...")

    application.run_polling()


if __name__ == "__main__":
    main()
